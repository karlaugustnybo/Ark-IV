"""Descriptive dated-tree prototype; never an Ark priority multiplier.

Scores use Duke's pinned BSD-3 engine. Identifier joins use only reviewed
one-to-one IUCN -> NCBI -> OTT links. A single median tree cannot establish
EDGE2 list membership across phylogenetic uncertainty.
"""

from __future__ import annotations

import csv
import hashlib
import io
import math
import re
import shutil
import tarfile
import urllib.request
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

import duckdb
import numpy as np
import pandas as pd

from ark_pipeline.runtime.provenance import iso_now, sha256

CATEGORIES = {"LC", "NT", "VU", "EN", "CR", "DD", "NE", "EX", "EW"}
GE = {"LC": 0, "NT": 1, "VU": 2, "EN": 3, "CR": 4}


def normalize_category(value: str | None) -> str:
    category = (value or "NE").strip().upper()
    category = {
        "LR/LC": "LC",
        "LR/NT": "NT",
        "LR/CD": "NT",
        "LEAST CONCERN": "LC",
        "NEAR THREATENED": "NT",
        "VULNERABLE": "VU",
        "ENDANGERED": "EN",
        "CRITICALLY ENDANGERED": "CR",
        "DATA DEFICIENT": "DD",
        "NOT EVALUATED": "NE",
        "EXTINCT": "EX",
        "EXTINCT IN THE WILD": "EW",
        "LOWER RISK/LEAST CONCERN": "LC",
        "LOWER RISK/NEAR THREATENED": "NT",
        "LOWER RISK/CONSERVATION DEPENDENT": "NT",
    }.get(category, category)
    if category not in CATEGORIES:
        raise ValueError(f"Unknown Red List category: {value!r}")
    return category


def original_edge(ed: float, category: str) -> float | None:
    """Isaac et al. 2007; DD/NE/EX/EW have no ordinary GE encoding."""
    return math.log1p(ed) + GE[category] * math.log(2) if category in GE else None


def acquire_inputs(config: dict, directory: Path, download: bool) -> dict:
    directory.mkdir(parents=True, exist_ok=True)
    records = {}
    for source in config["sources"]:
        path = directory / source["filename"]
        if not path.is_file():
            if not download:
                raise ValueError(f"Missing {path}; run just edge-download")
            print(f"Downloading {source['id']} ({source['release']})", flush=True)
            temporary = path.with_suffix(path.suffix + ".part")
            with (
                urllib.request.urlopen(source["url"], timeout=120) as response,
                temporary.open("wb") as output,
            ):
                shutil.copyfileobj(response, output)
            verify_source(temporary, source)
            temporary.replace(path)
        verify_source(path, source)
        records[source["id"]] = {
            **source,
            "path": str(path.resolve()),
            "bytes": path.stat().st_size,
            "sha256": sha256(path),
            "verified_at": iso_now(),
        }
    return records


def verify_source(path: Path, source: dict) -> None:
    for algorithm in ("sha256", "md5"):
        if algorithm in source:
            with path.open("rb") as handle:
                digest = hashlib.file_digest(handle, algorithm).hexdigest()
            if digest != source[algorithm]:
                raise ValueError(f"{source['id']} {algorithm} checksum changed: {path}")


def extract_clade(newick: str, label: str) -> str:
    """Extract a named crown without materializing 2.3 million ETE nodes.

    The released tree uses unquoted labels. Reject quoted/commented syntax
    rather than interpreting parentheses inside a quoted taxon name.
    """
    matches = list(re.finditer(r"\)" + re.escape(label) + r"(?=[:,;)])", newick))
    if len(matches) != 1:
        raise ValueError(f"Expected one internal clade {label}, found {len(matches)}")
    end = matches[0].start()
    depth = 1
    start = end
    while start > 0 and depth:
        start -= 1
        char = newick[start]
        if char in "'\"[]":
            raise ValueError("Clade extraction requires unquoted, uncommented Newick labels")
        depth += (char == ")") - (char == "(")
    if depth:
        raise ValueError("Unbalanced clade Newick")
    # A crown-based comparison excludes the mammalian stem branch.
    return newick[start : end + 1] + label + ":0;"


def load_clade(path: Path, label: str):
    from ete4 import Tree

    tree = Tree(extract_clade(path.read_text(), label), parser=1)
    for node in tree.traverse():
        if node.dist is None or not math.isfinite(node.dist) or node.dist < 0:
            raise ValueError(f"Invalid branch length on {node.name}")
    labels = [node.name for node in tree.leaves()]
    if len(labels) != len(set(labels)):
        raise ValueError("Duplicate tree tip labels")
    tree.dist = 0
    return tree


def load_ncbi_merges(path: Path | None) -> dict[str, str]:
    if path is None:
        return {}
    with path.open() as handle:
        return {
            parts[0].strip(): parts[1].strip()
            for line in handle
            if len(parts := line.split("|")) >= 2
        }


def canonical_taxid(taxid: str, merges: dict[str, str]) -> str:
    seen = set()
    while taxid in merges:
        if taxid in seen:
            raise ValueError("Cycle in NCBI merged IDs")
        seen.add(taxid)
        taxid = merges[taxid]
    return taxid


def tip_identifiers(tree, taxonomy_archive: Path, merges: dict[str, str]) -> dict[str, set[str]]:
    """Stream the OTT taxonomy; retain identifiers for complete-clade tips only."""
    tips = {}
    for leaf in tree.leaves():
        match = re.search(r"_ott(\d+)$", leaf.name)
        if match is None:
            raise ValueError(f"Tip lacks OTT identifier: {leaf.name}")
        ottid = match[1]
        # The release includes imputed tips with the placeholder ott0000.
        # They remain part of the complete clade, without an identifier join.
        if int(ottid) == 0:
            continue
        if ottid in tips:
            raise ValueError(f"Duplicate OTT tip identifier: {ottid}")
        tips[ottid] = leaf.name
    identifiers: dict[str, set[str]] = defaultdict(set)
    taxonomy_tips = {}
    with tarfile.open(taxonomy_archive) as archive:
        members = [item for item in archive.getmembers() if item.name.endswith("/taxonomy.tsv")]
        if len(members) != 1:
            raise ValueError("OTT archive must contain exactly one taxonomy.tsv")
        with io.TextIOWrapper(archive.extractfile(members[0]), encoding="utf-8") as handle:
            header = [part.strip() for part in next(handle).split("|")]
            uid_index, source_index = header.index("uid"), header.index("sourceinfo")
            rank_index = header.index("rank")
            for line in handle:
                parts = [part.strip() for part in line.split("|")]
                if parts[uid_index] in tips:
                    taxonomy_tips[tips[parts[uid_index]]] = {
                        "rank": parts[rank_index],
                        "flags": parts[header.index("flags")],
                    }
                if parts[uid_index] in tips and parts[rank_index] == "species":
                    for source in parts[source_index].split(","):
                        if source.startswith("ncbi:"):
                            taxid = canonical_taxid(source.removeprefix("ncbi:"), merges)
                            identifiers[taxid].add(tips[parts[uid_index]])
    for leaf in tree.leaves():
        metadata = taxonomy_tips.get(leaf.name, {})
        leaf.add_prop("ott_rank", metadata.get("rank"))
        leaf.add_prop("ott_flags", metadata.get("flags"))
    return dict(identifiers)


def assessment_join(
    crosswalk: Path, assessments: Path, clade_class: str | None = "MAMMALIA"
) -> list[dict]:
    with duckdb.connect() as connection:
        connection.execute(
            "CREATE TEMP TABLE cw AS SELECT * FROM read_parquet(?)", [str(crosswalk)]
        )
        required = {
            "iucn_sis_id",
            "iucn_assessment_id",
            "iucn_scientific_name",
            "iucn_class",
            "matched_ncbi_species_taxid",
            "safe_for_automatic_species_trait_transfer",
            "match_status",
        }
        columns = {row[0] for row in connection.execute("DESCRIBE cw").fetchall()}
        if not required.issubset(columns):
            raise ValueError(f"Crosswalk missing columns: {sorted(required - columns)}")
        connection.execute(
            "CREATE TEMP TABLE a AS SELECT * FROM read_csv(?, all_varchar=true)", [str(assessments)]
        )
        if connection.execute(
            "SELECT count(*) FROM (SELECT assessmentId FROM a GROUP BY ALL HAVING count(*)>1)"
        ).fetchone()[0]:
            raise ValueError("Duplicate assessment IDs")
        result = connection.execute(
            """
            SELECT cw.iucn_sis_id, cw.iucn_assessment_id, cw.iucn_scientific_name, cw.iucn_class,
                   cw.matched_ncbi_species_taxid, cw.safe_for_automatic_species_trait_transfer,
                   cw.match_status, a.redlistCategory, a.scopes, a.possiblyExtinct,
                   a.possiblyExtinctInTheWild, a.yearPublished, a.assessmentDate,
                   a.internalTaxonId AS assessment_sis_id
            FROM cw LEFT JOIN a ON cast(cw.iucn_assessment_id AS VARCHAR)=a.assessmentId
            WHERE (? IS NULL OR upper(cw.iucn_class)=?)
            ORDER BY cw.iucn_sis_id
        """,
            [clade_class, clade_class],
        ).fetchdf()
    # Convert SQL nulls to Python None without making boolean nulls truthy.
    rows = result.astype(object).where(pd.notna(result), None).to_dict("records")
    ids = [row["iucn_sis_id"] for row in rows]
    if len(ids) != len(set(ids)):
        raise ValueError("Duplicate IUCN SIS IDs in crosswalk")
    return rows


def reviewed_tip_join(
    rows: list[dict], identifiers: dict[str, set[str]], merges: dict[str, str]
) -> list[dict]:
    candidates = Counter()
    joined = []
    for original in rows:
        row = dict(original)
        row.update(tree_tip=None, join_status="unsafe_crosswalk", redlist_category=None)
        if row["redlistCategory"] is not None:
            row["redlist_category"] = normalize_category(row["redlistCategory"])
        if row["safe_for_automatic_species_trait_transfer"] is True:
            if row["assessment_sis_id"] is None:
                row["join_status"] = "assessment_not_found"
            elif str(row["assessment_sis_id"]) != str(row["iucn_sis_id"]):
                row["join_status"] = "assessment_identifier_conflict"
            elif (
                not re.search(r"\bGlobal\b", row["scopes"] or "", re.I)
                or "subpopulation" in row["iucn_scientific_name"].lower()
            ):
                row["join_status"] = "non_global_or_subpopulation"
            elif not row["matched_ncbi_species_taxid"]:
                row["join_status"] = "no_ncbi_id"
            else:
                taxid = canonical_taxid(str(row["matched_ncbi_species_taxid"]), merges)
                tips = identifiers.get(taxid, set())
                row["canonical_ncbi_taxid"] = taxid
                row["join_status"] = "ncbi_not_in_clade" if not tips else "ambiguous_ncbi_to_ott"
                if len(tips) == 1:
                    row["tree_tip"] = next(iter(tips))
                    row["join_status"] = "joined"
                    candidates[row["tree_tip"]] += 1
        joined.append(row)
    for row in joined:
        if row["join_status"] == "joined" and candidates[row["tree_tip"]] != 1:
            row["join_status"] = "ambiguous_iucn_to_tip"
    return joined


def candidate_list(category: str, above_fraction: float) -> str | None:
    if above_fraction < 0.95:
        return None
    if category in {"CR", "EN", "VU"}:
        return "threatened"
    if category in {"DD", "NE"}:
        return "research"
    if category in {"LC", "NT"}:
        return "watch"
    return None


def score_clade(
    tree, joined: list[dict], samples: int, seed: int, *, stable_risk_order: bool = False
) -> tuple[list[dict], dict]:
    from dated_complete_tree import tree_metrics as engine
    from dated_complete_tree.tree_fixing_utils import open_config_csv

    from ark_pipeline.extinction_status import extinction_status

    if samples < 2:
        raise ValueError("Use at least two risk samples")
    by_tip = {row["tree_tip"]: row for row in joined if row["join_status"] == "joined"}
    status_by_tip = {
        leaf.name: extinction_status(
            by_tip.get(leaf.name, {}).get("redlist_category"), leaf.props.get("ott_flags")
        )
        for leaf in tree.leaves()
    }
    # Already extinct species do not have avertable evolutionary loss.
    retained = [
        leaf.name
        for leaf in tree.leaves()
        if status_by_tip[leaf.name]["effective_extinct"] is not True
    ]
    original_count = len(tree)
    if len(retained) < 2:
        raise ValueError("Fewer than two extant tips")
    if len(retained) != original_count:
        tree.prune(retained, preserve_branch_length=True)
    tree.dist = 0
    leaves = list(tree.leaves())
    for leaf in leaves:
        category = by_tip.get(leaf.name, {}).get("redlist_category") or "NE"
        leaf.add_prop("iucn_status", category)
    engine.label_desc_leaves(tree)
    engine.compute_ed_scores(tree)
    with open_config_csv(None, "p_extinction.csv") as handle:
        risk_rows = list(csv.reader(handle))[1:]
    probabilities = [float(row[1]) for row in risk_rows]
    risk_lookup = {
        category: [float(row[1]) for row in risk_rows if row[0] == category]
        for category in {row[0] for row in risk_rows}
    }
    stable_leaves = sorted(leaves, key=lambda leaf: leaf.name)
    missing_median = float(np.median(probabilities))
    rng = np.random.default_rng(seed)

    def compute(randomise: bool):
        if stable_risk_order:
            # Identical tip sets receive identical draws across alternative trees,
            # regardless of their Newick traversal order. Tree effects then do not
            # include changes in which species received which random draw.
            uniforms = rng.random(len(stable_leaves)) if randomise else None
            for risk_index, leaf in enumerate(stable_leaves):
                values = risk_lookup.get(leaf.props["iucn_status"], probabilities)
                risk = (
                    values[int(uniforms[risk_index] * len(values))]
                    if randomise
                    else np.median(values)
                )
                leaf.add_prop("pext", float(risk))
        else:
            engine.assign_extinction_risks(
                tree,
                rng=rng,
                randomise_risk=randomise,
                missing_value=None if randomise else missing_median,
            )
        for leaf in leaves:
            if leaf.props["iucn_status"] == "EW":
                leaf.add_prop("pext", 0.97)
        engine.compute_edge2_scores(tree)
        return np.array(
            [[leaf.props[key] for key in ("ed2_score", "pext", "edge2_score")] for leaf in leaves],
            dtype=float,
        )

    point = compute(False)
    distributions = np.empty((samples, len(leaves), 3))
    above = np.zeros(len(leaves), dtype=int)
    for index in range(samples):
        values = compute(True)
        distributions[index] = values
        above += values[:, 2] > np.median(values[:, 2])
    if not np.isfinite(distributions).all() or np.any(distributions < 0):
        raise ValueError("Engine returned invalid scores")
    summaries = np.quantile(distributions, [0.025, 0.25, 0.5, 0.75, 0.975], axis=0)
    output = []
    for index, leaf in enumerate(leaves):
        category = leaf.props["iucn_status"]
        record = {
            "tree_tip": leaf.name,
            "redlist_category": category,
            "ott_rank": leaf.props.get("ott_rank"),
            "ott_flags": leaf.props.get("ott_flags"),
            **status_by_tip[leaf.name],
            "ark_ed_myr": float(leaf.props["ed_score"]),
            "terminal_branch_myr": float(leaf.dist),
            "ark_edge1": original_edge(leaf.props["ed_score"], category),
            "above_clade_median_fraction": float(above[index] / samples),
            "risk_only_candidate_list": candidate_list(category, above[index] / samples),
            "risk_samples": samples,
        }
        for component, column in enumerate(("ark_ed2_myr", "ark_ge2", "ark_edge2_myr")):
            record[column + "_point"] = float(point[index, component])
            for qi, quantile in enumerate(("p025", "p25", "median", "p75", "p975")):
                record[column + "_" + quantile] = float(summaries[qi, index, component])
        output.append(record)
    diagnostics = {
        "original_tree_tips": original_count,
        "retained_tree_tips": len(leaves),
        "excluded_extinct_tips": original_count - len(leaves),
        "unassessed_or_unjoined_tree_tips": sum(leaf.name not in by_tip for leaf in leaves),
        "sum_ed_myr": sum(row["ark_ed_myr"] for row in output),
        "clade_pd_myr": float(engine.compute_pd(tree)),
        "unknown_risk_point": missing_median,
        "risk_draw_order": "sorted_tip_name" if stable_risk_order else "tree_traversal",
    }
    diagnostics["ott_tip_ranks"] = dict(
        Counter(leaf.props.get("ott_rank") or "placeholder_or_missing" for leaf in leaves)
    )
    diagnostics["ott_extinct_flag_tips"] = [
        leaf.name for leaf in leaves if "extinct" in (leaf.props.get("ott_flags") or "")
    ]
    diagnostics["extinction_status_policy"] = "verified_iucn_overrides_ott_extinction_flags"
    diagnostics["iucn_overrides_ott_extinction_tips"] = [
        name for name, status in status_by_tip.items() if status["ott_extinction_overridden"]
    ]
    if not math.isclose(diagnostics["sum_ed_myr"], diagnostics["clade_pd_myr"], rel_tol=1e-10):
        raise ValueError("ED does not sum to complete-clade phylogenetic diversity")
    return output, diagnostics


def read_mammal_reference(path: Path) -> list[dict]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["All EDGE2 Scores"]
        values = sheet.iter_rows(values_only=True)
        header = next(values)
        rows = [dict(zip(header, row, strict=True)) for row in values if row[0] is not None]
        required = {"RL.ID", "Species", "RL.cat", "ED.median", "EDGE.median"}
        if not required.issubset(header):
            raise ValueError("Mammal reference schema changed")
        result = []
        for row in rows:
            # The supplement's RL.ID is the stable SIS taxon ID, not assessmentId.
            raw_sis = str(row["RL.ID"]).strip()
            if raw_sis not in {"NA", "None", ""} and not raw_sis.isdigit():
                raise ValueError(f"Unexpected reference SIS identifier: {raw_sis}")
            result.append(
                {
                    "iucn_sis_id": int(raw_sis) if raw_sis.isdigit() else None,
                    "reference_name": row["Species"],
                    "reference_category": normalize_category(
                        None if row["RL.cat"] == "NA" else row["RL.cat"]
                    ),
                    "reference_ed2_myr": float(row["ED.median"]),
                    "reference_edge2_myr": float(row["EDGE.median"]),
                    "reference_terminal_myr": float(row["TBL.median"]),
                }
            )
        return result
    finally:
        workbook.close()


def compare_reference(species: pd.DataFrame, reference: list[dict]) -> tuple[pd.DataFrame, dict]:
    ref = pd.DataFrame(reference)
    # Some unassessed reference records use 0/placeholder IDs. They cannot be
    # joined to SIS records; duplicate real IDs also cannot be selected arbitrarily.
    usable = ref[(ref.iucn_sis_id > 0) & ~ref.iucn_sis_id.duplicated(keep=False)]
    comparison = species[species.join_status == "scored"].merge(
        usable, on="iucn_sis_id", validate="one_to_one"
    )
    comparison["category_agrees"] = comparison.redlist_category == comparison.reference_category
    for component in ("ed2", "edge2"):
        comparison[component + "_residual_myr"] = (
            comparison[f"ark_{component}_myr_median"] - comparison[f"reference_{component}_myr"]
        )
        comparison[component + "_rank_delta"] = comparison[f"ark_{component}_myr_median"].rank(
            ascending=False
        ) - comparison[f"reference_{component}_myr"].rank(ascending=False)

    def statistics(frame: pd.DataFrame) -> dict:
        result: dict[str, Any] = {"species": len(frame)}
        for component in ("ed2", "edge2"):
            x, y = frame[f"ark_{component}_myr_median"], frame[f"reference_{component}_myr"]
            correlation = (
                x.rank().corr(y.rank()) if x.nunique() > 1 and y.nunique() > 1 else float("nan")
            )
            result[component + "_spearman"] = float(correlation) if pd.notna(correlation) else None
            result[component + "_median_abs_residual_myr"] = (
                float((x - y).abs().median()) if len(frame) else None
            )
        return result

    return comparison, {
        "reference_rows": len(ref),
        "unique_positive_sis_ids": len(usable),
        "reference_without_sis_ids": int(ref.iucn_sis_id.isna().sum()),
        "all_joined": statistics(comparison),
        "same_category": statistics(comparison[comparison.category_agrees]),
        "category_changed_species": int((~comparison.category_agrees).sum()),
        "interpretation": "Sensitivity comparison: different tree, assessment snapshot, and risk-only sampling; not reproduction of the published 1,000-tree result.",
    }
